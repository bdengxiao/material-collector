"""Offline collection engine. Never mutates source materials."""
from __future__ import annotations
import csv, hashlib, json, os, re, shutil, tempfile, uuid, zipfile, unicodedata
from pathlib import Path
from datetime import datetime, timezone
from dataclasses import dataclass, field, asdict
from functools import lru_cache

def now(): return datetime.now(timezone.utc).isoformat()
def digest(path):
    h=hashlib.sha256()
    with open(path,'rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()

@dataclass
class Person:
    key: str
    name: str
    identifier: str=''

@dataclass
class Rule:
    name: str
    extensions: list[str]
    required: bool=True
    multiple: bool=False
    keywords: list[str]=field(default_factory=list)

@dataclass
class Material:
    path: str
    relative: str
    size: int
    sha256: str
    person: str|None=None
    category: str|None=None
    reason: str='待确认'
    manual: bool=False
    ignored: bool=False
    mtime_ns: int=0
    invalid: str=''
    needs_review: bool=False
    file_id: str=''
    previous_person: str|None=None
    previous_category: str|None=None

@dataclass
class Project:
    title: str='新核对任务'
    project_id: str=field(default_factory=lambda:uuid.uuid4().hex)
    people: list[Person]=field(default_factory=list)
    rules: list[Rule]=field(default_factory=list)
    materials: list[Material]=field(default_factory=list)
    roots: list[str]=field(default_factory=list)
    events: list[dict]=field(default_factory=list)
    exports: list[dict]=field(default_factory=list)
    def event(self,action,**details): self.events.append(dict(time=now(),action=action,**details))

TEMPLATES={
 '活动报名材料':[Rule('报名表',['.docx','.pdf'],keywords=['报名']),Rule('照片',['.jpg','.jpeg','.png'])],
 '社团作品收集':[Rule('作品',['.pdf'],keywords=['作品']),Rule('说明',['.docx','.txt'],keywords=['说明'])],
 '班级电子资料核对':[Rule('证明',['.pdf','.jpg','.png'],keywords=['证明']),Rule('登记表',['.xlsx'],keywords=['登记'])]}

def read_roster(path):
    path=Path(path); warnings=[]
    if path.suffix.lower()=='.csv':
        raw=path.read_bytes()
        for enc in ('utf-8-sig','gb18030','utf-16'):
            try: text=raw.decode(enc); break
            except UnicodeError: continue
        else: raise ValueError('名单编码无法识别，请另存为 UTF-8 CSV。')
        rows=list(csv.reader(text.splitlines()))
    elif path.suffix.lower()=='.xlsx':
        from openpyxl import load_workbook
        try:wb=load_workbook(path,read_only=True,data_only=True)
        except Exception as e:raise ValueError('Excel 名单无法读取，可能损坏、加密或不是真正的 XLSX。请在 Excel 中另存为 XLSX 后重试。') from e
        try:
            rows=[]
            for cells in wb.active.iter_rows():
                row=[]
                for c in cells:
                    v=c.value
                    if isinstance(v,(int,float)):
                        if abs(v)>=10**15:warnings.append('编号超过15位，补零格式不能恢复精度，请使用原始文本编号。')
                        if isinstance(v,float) and not v.is_integer(): warnings.append('存在小数数字，请确认编号列。')
                        fmt=c.number_format
                        if re.fullmatch('0+',fmt or '') and isinstance(v,int): v=str(v).zfill(len(fmt))
                        else:
                            warnings.append('存在数字格式单元格；如编号原有前导零但未设置补零格式，本软件无法还原。请核对原始文本编号。')
                            if abs(v)>=10**15: warnings.append('数字达到15位以上，Excel可能已丢失精度；请使用原始文本编号。')
                            v=str(int(v)) if float(v).is_integer() else str(v)
                    row.append('' if v is None else str(v).strip())
                rows.append(row)
        finally: wb.close()
    else: raise ValueError('请选择 .xlsx 或 .csv 名单。')
    rows=[r for r in rows if any(str(x).strip() for x in r)]
    if len(rows)<2: raise ValueError('名单没有数据行。')
    headers=[str(x).strip() for x in rows[0]]
    if not any(headers): raise ValueError('第一行需要表头。')
    return headers,rows[1:],list(dict.fromkeys(warnings))

def parse_people(headers,rows,name_col,id_col=None):
    if name_col<0 or name_col>=len(headers): raise ValueError('姓名列无效。')
    warnings=[]; people=[]
    if id_col is not None and any(x in headers[id_col] for x in ('后四','后4','手机','电话')):
        warnings.append('手机号等字段不作为自动匹配编号。'); id_col=None
    for r in rows:
        name=str(r[name_col]).strip() if name_col<len(r) else ''
        ident=str(r[id_col]).strip() if id_col is not None and id_col<len(r) else ''
        if name: people.append(Person(uuid.uuid4().hex,name,ident))
    if not people: raise ValueError('所选姓名列没有有效姓名。')
    for attr,label in [('name','姓名'),('identifier','编号')]:
        values=[getattr(p,attr) for p in people if getattr(p,attr)]
        duplicates=sorted({v for v in values if values.count(v)>1})
        if duplicates: warnings.append(f'重复{label}：'+ '、'.join(duplicates[:20]))
    return people,warnings

def linked(path):
    return path.is_symlink() or bool(getattr(path.stat(follow_symlinks=False),'st_file_attributes',0)&0x400)

def reject_link_path(path):
    value=os.fspath(path)
    if os.name=='nt' and (':' in value[2:] or value.startswith(('\\\\?\\','\\\\.\\'))):raise ValueError('不支持备用数据流或设备命名空间路径')
    absolute=Path(os.path.abspath(path))
    for item in [absolute,*absolute.parents]:
        if item.exists() and linked(item): raise ValueError('路径不能经过链接或目录联接：'+str(item))
    return absolute

def scan(root,cancel=lambda:False,progress=lambda x:None):
    root=reject_link_path(root).resolve()
    if not root.is_dir() or linked(root): raise ValueError('请选择正常材料文件夹，不能选择链接目录。')
    found=[]; errors=[]
    def fail(e): errors.append(str(e))
    for directory,dirs,files in os.walk(root,followlinks=False,onerror=fail):
        safe=[]
        for d in dirs:
            try:
                if d not in ('$RECYCLE.BIN','System Volume Information','.git') and not linked(Path(directory)/d): safe.append(d)
            except OSError as e: errors.append(str(e))
        dirs[:]=safe
        for name in files:
            if cancel(): raise InterruptedError('扫描已取消，原项目未更新。')
            p=Path(directory)/name
            try:
                if linked(p): errors.append(f'跳过链接：{p}'); continue
                stat=p.stat(); h=digest(p)
                if p.stat().st_mtime_ns!=stat.st_mtime_ns: raise OSError('扫描时文件发生变化')
                found.append(Material(str(p),str(p.relative_to(root)),stat.st_size,h,mtime_ns=stat.st_mtime_ns,file_id=f'{stat.st_dev}:{stat.st_ino}')); progress(len(found))
            except OSError as e: errors.append(f'{p}: {e}')
    return found,errors

@lru_cache(maxsize=20000)
def identifier_pattern(value):return re.compile(r'(?<![^\W_])'+re.escape(value)+r'(?![^\W_])',re.I)

def token(text,value):
    if not value:return False
    # Marks are part of an identity token, never reliable separators.
    for hit in identifier_pattern(value).finditer(text):
        adjacent=(text[hit.start()-1:hit.start()] if hit.start() else '')+text[hit.end():hit.end()+1]
        if not any(unicodedata.category(c)[0] in 'LMN' for c in adjacent):return True
    return False

def match(project,m):
    if m.manual or m.ignored or m.needs_review: return
    m.person=None; m.category=None
    text=Path(m.relative).stem
    ids=[p for p in project.people if token(text,p.identifier)]
    if ids:
        conflicting_names=[p for p in project.people if p.name in text and p.key!=ids[0].key and p.name!=ids[0].name] if len(ids)==1 else []
        if len(ids)==1 and not conflicting_names: m.person=ids[0].key; m.reason='完整唯一编号'
        elif conflicting_names: m.reason='编号与姓名矛盾，待确认'
        else: m.reason='编号冲突，待确认'
    else:
        # Longer names shadow their shorter substrings; repeated names never auto-assign.
        hits=[p for p in project.people if p.name and p.name in text]
        hits=[p for p in hits if not any(p.name!=q.name and p.name in q.name for q in hits)]
        if len(hits)==1 and sum(p.name==hits[0].name for p in project.people)==1:
            # Chinese adjacent characters make an unbounded name uncertain.
            if token(text,hits[0].name):
                m.person=hits[0].key; m.reason='唯一姓名，完整分隔边界'
            else: m.reason='姓名边界不明确，待确认'
        else: m.reason='姓名重名或无法识别，待确认'
    ext=Path(m.path).suffix.lower()
    eligible=[r for r in project.rules if ext in r.extensions]
    strong=[r for r in eligible if r.keywords and any(k in text for k in r.keywords)]
    chosen=strong if strong else eligible
    if len(chosen)==1: m.category=chosen[0].name
    elif not chosen: m.reason+='；文件格式不符合规则'
    else: m.reason+='；材料类别待确认'

def merge_scan(project,root,found):
    old={m.path:m for m in project.materials}; merged=dict(old)
    scanned_root=Path(root).resolve(); found_paths={m.path for m in found}
    for path,m in merged.items():
        if Path(path).is_relative_to(scanned_root) and path not in found_paths:
            m.invalid='本次扫描未找到或无法访问';m.reason=m.invalid+'；重新扫描原目录或在人工页明确忽略'
    for m in found:
        prior=old.get(m.path)
        if prior and prior.sha256==m.sha256 and (not (prior.manual or prior.ignored) or (prior.file_id and prior.file_id==m.file_id and prior.mtime_ns==m.mtime_ns)):
            m.person,m.category,m.manual,m.ignored,m.reason=prior.person,prior.category,prior.manual,prior.ignored,prior.reason
            m.needs_review=prior.needs_review
            m.previous_person,m.previous_category=prior.previous_person,prior.previous_category
        elif prior:
            m.previous_person,m.previous_category=prior.person or prior.previous_person,prior.category or prior.previous_category
            m.reason='内容或文件身份变化，旧确认仅作参考，请重新确认';m.needs_review=True;m.manual=False;m.ignored=False
        match(project,m); merged[m.path]=m
    project.materials=list(merged.values())
    if str(Path(root).resolve()) not in project.roots: project.roots.append(str(Path(root).resolve()))
    project.event('导入材料',root=str(root),count=len(found))

def unchanged(m):
    try:
        p=Path(m.path);s=p.stat()
        return not linked(p) and s.st_size==m.size and s.st_mtime_ns==m.mtime_ns and (not m.file_id or m.file_id==f'{s.st_dev}:{s.st_ino}')
    except OSError:return False

def status(project,p):
    mats=[m for m in project.materials if m.person==p.key and not m.ignored and not m.invalid and not m.needs_review and m.category and unchanged(m) and Path(m.path).suffix.lower() in next((r.extensions for r in project.rules if r.name==m.category),[])]
    missing=[r.name for r in project.rules if r.required and not any(m.category==r.name for m in mats)]
    duplicate=[r.name for r in project.rules if not r.multiple and sum(m.category==r.name for m in mats)>1]
    return mats,missing,duplicate

def assign(project,index,person,category,ignored=False):
    if person and person not in {p.key for p in project.people}: raise ValueError('名单人员无效')
    if category and category not in {r.name for r in project.rules}: raise ValueError('材料类别无效')
    m=project.materials[index]
    if not unchanged(m):raise ValueError('源文件变化或无法访问，请重新扫描原目录后确认。丢失文件可明确忽略，缺件仍会保留。')
    if digest(reject_link_path(m.path))!=m.sha256:raise ValueError('源文件内容变化，请先重新扫描后确认。')
    if category and Path(m.path).suffix.lower() not in next(r.extensions for r in project.rules if r.name==category): raise ValueError('文件格式不符合所选材料规则，请修改规则或选择正确类别。')
    m.person=person; m.category=category; m.ignored=ignored; m.manual=True;m.invalid='';m.needs_review=False; m.reason='人工确认' if person else '人工清除归属'
    st=Path(m.path).stat();m.file_id=f'{st.st_dev}:{st.st_ino}';m.previous_person=None;m.previous_category=None
    project.event('人工更正',path=m.path,person=person,category=category,ignored=ignored)

def update_rules(project,rules):
    old={r.name:asdict(r) for r in project.rules};new={r.name:asdict(r) for r in rules};affected=0
    for m in project.materials:
        if m.manual:
            if m.category and old.get(m.category)!=new.get(m.category):
                m.needs_review=True;m.reason='材料规则变化，原人工归属保留供参考，请重新确认';affected+=1
        else:match(Project(people=project.people,rules=rules),m)
    project.rules=rules;project.event('更新规则',affected_manual=affected);return affected

def save(project,path):
    path=reject_link_path(path); path.parent.mkdir(parents=True,exist_ok=True)
    if path.exists():
        try:
            existing=load(path)
            if existing.project_id!=project.project_id:raise ValueError('已有文件属于另一个项目，请另选新名称')
        except Exception as e:raise ValueError('已有文件不是有效的核对项目，拒绝覆盖。请另选新文件名。') from e
    for m in project.materials:
        source=Path(m.path)
        if path==source or (path.exists() and source.exists() and os.path.samefile(path,source)): raise ValueError('项目不能保存在原材料文件上。')
    backup=Path(str(path)+'.bak')
    if backup.exists():
        reject_link_path(backup)
        try:
            old_backup=load(backup)
            if not path.exists() or old_backup.project_id!=load(path).project_id:raise ValueError('备份属于另一个项目')
        except Exception as e:raise ValueError('备份位置已有无关或损坏文件，拒绝覆盖。请另选新项目文件名。') from e
        if backup.stat().st_nlink>1: raise ValueError('项目备份是硬链接，拒绝写入。')
        for m in project.materials:
            if Path(m.path).exists() and os.path.samefile(backup,m.path): raise ValueError('备份路径与原材料冲突。')
    data={'schema':1,'saved_at':now(),'project':asdict(project)}
    fd,temp=tempfile.mkstemp(dir=path.parent,prefix='.save-')
    try:
        with os.fdopen(fd,'w',encoding='utf-8') as f: json.dump(data,f,ensure_ascii=False,indent=2); f.flush(); os.fsync(f.fileno())
        if path.exists():
            fd2,temp2=tempfile.mkstemp(dir=path.parent,prefix='.backup-')
            try:
                with os.fdopen(fd2,'wb') as out,path.open('rb') as source:shutil.copyfileobj(source,out);out.flush();os.fsync(out.fileno())
                archive=Path(str(path)+'.history-'+uuid.uuid4().hex+'.msq');os.replace(temp2,archive)
                shutil.copy2(archive,temp2)
                os.replace(temp,path)
                os.replace(temp2,backup)
            finally:
                if os.path.exists(temp2):os.unlink(temp2)
        if os.path.exists(temp):os.replace(temp,path)
    finally:
        if os.path.exists(temp): os.unlink(temp)

def load(path):
    d=json.loads(Path(path).read_text(encoding='utf-8-sig'))
    if d.get('schema')!=1: raise ValueError('项目版本不支持')
    p=d['project']; p['people']=[Person(**x) for x in p['people']]; p['rules']=[Rule(**x) for x in p['rules']]; p['materials']=[Material(**x) for x in p['materials']]
    obj=Project(**p)
    for m in obj.materials:
        reject_link_path(m.path)
        if not Path(m.path).is_absolute():raise ValueError('材料源路径必须是绝对路径')
        if (m.manual or m.ignored) and not m.file_id:m.needs_review=True;m.ignored=False;m.reason='旧项目缺少文件身份快照，请重新扫描确认'
    if len({x.key for x in obj.people})!=len(obj.people): raise ValueError('项目人员标识重复')
    return obj

def safe_name(s):
    s=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',s).strip(' .')[:80] or '未命名'
    if s.split('.')[0].upper() in {'CON','PRN','AUX','NUL',*[f'COM{i}' for i in range(1,10)],*[f'LPT{i}' for i in range(1,10)]}: s='_'+s
    return s

def export_plan(project,parent):
    parent=reject_link_path(parent).resolve()
    for root in project.roots:
        r=Path(root).resolve()
        if parent==r or parent.is_relative_to(r) or r.is_relative_to(parent): raise ValueError('输出目录与材料目录不能互相包含，请选择独立目录。')
    for ancestor in [parent,*parent.parents]:
        if ancestor.exists() and linked(ancestor): raise ValueError('输出路径不能经过链接目录')
    people={p.key:p for p in project.people}; plan=[]
    for i,m in enumerate(project.materials):
        if m.ignored: continue
        if m.invalid or m.needs_review or not unchanged(m):raise ValueError(f'导出前需处理失效/待复核文件：{m.relative}。请重新扫描并确认；文件丢失可在人工页明确忽略，报告仍显示缺件。')
        p=people.get(m.person)
        directory=(safe_name(p.name+'_'+p.identifier)+'_'+safe_name(p.key[:8])) if p else '待确认文件'
        name=safe_name(m.category or '未分类')+'_'+str(i+1)+Path(m.path).suffix.lower()
        plan.append((m,Path('材料')/directory/name))
    return plan

def export(project,parent,cancel=lambda:False,progress=lambda x:None):
    if cancel():raise InterruptedError('导出已取消，未创建输出。')
    from openpyxl import Workbook,load_workbook
    plan=export_plan(project,parent); parent=Path(parent).resolve(); parent.mkdir(parents=True,exist_ok=True)
    estimated=sum(m.size for m,_ in plan)*2+10*1024*1024
    if shutil.disk_usage(parent).free<estimated:raise OSError('输出空间不足，请选择有足够剩余空间的磁盘。')
    target=parent/('整理结果_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:8]); target.mkdir()
    state={'project_id':project.project_id,'generated_at':now(),'status':'进行中','files':[]}
    def manifest(): (target/'导出记录.json').write_text(json.dumps(state,ensure_ascii=False,indent=2),encoding='utf-8')
    def sheet(filename,headers,rows):
        if cancel():raise InterruptedError('报告导出已取消，结果未完成。')
        def cell(v):return ("'"+v) if isinstance(v,str) and v.startswith(('=','+','-','@')) else v
        wb=Workbook(); ws=wb.active; ws.title='核对记录'; ws.append([cell(v) for v in headers])
        for row in rows:
            ws.append([("'"+v) if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in row])
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=24
        ws.append([cell(v) for v in ['项目标识',project.project_id,'生成时间',state['generated_at']]]); wb.save(target/filename)
        check=load_workbook(target/filename,read_only=True); check.close()
    try:
        manifest()
        for i,(m,relative) in enumerate(plan):
            if cancel(): raise InterruptedError('整理已取消，部分结果保留并标记未完成。')
            src=reject_link_path(m.path)
            if linked(src) or digest(src)!=m.sha256: raise ValueError(f'源文件变化或链接，需要重新核对：{src}')
            dest=target/relative; dest.parent.mkdir(parents=True,exist_ok=True);reject_link_path(dest)
            with src.open('rb') as f,dest.open('xb') as out: shutil.copyfileobj(f,out)
            if digest(dest)!=m.sha256: raise OSError(f'复制校验失败：{src}')
            state['files'].append({'source':m.path,'output':str(relative),'sha256':m.sha256,'person':m.person,'category':m.category,'reason':m.reason}); progress(i+1)
        rows=[]; missing=[]
        for p in project.people:
            mats,miss,dup=status(project,p); row=[p.name,p.identifier,'、'.join(r.name for r in project.rules if r.required),'、'.join(sorted({m.category for m in mats if m.category})),'、'.join(miss),'重复：'+ '、'.join(dup) if dup else '', '文件已找到不代表内容合格']
            row.extend(('缺失' if r.name in miss else '重复' if r.name in dup else '已找到' if any(m.category==r.name for m in mats) else '选交未找到') for r in project.rules)
            rows.append(row)
            if miss: missing.append(row)
        headers=['姓名','编号','应交材料','已找到材料','缺失材料','异常','说明']+[r.name for r in project.rules]
        sheet('核对总表.xlsx',headers,rows); sheet('缺件清单.xlsx',headers,missing)
        anomalies=[]
        for m in project.materials:
            dup=False
            if m.person:
                p=next((p for p in project.people if p.key==m.person),None)
                if p: dup=m.category in status(project,p)[2]
            if not m.person or not m.category or dup or not unchanged(m): anomalies.append([m.relative,m.person or '',m.category or '',m.reason,'源文件变化/不可访问' if not unchanged(m) else '重复提交' if dup else '待确认'])
        sheet('异常文件清单.xlsx',['文件','人员标识','类别','依据','状态'],anomalies)
        state['status']='打包中'; manifest()
        zip_path=target/'整理结果.zip'
        with zipfile.ZipFile(zip_path,'x',zipfile.ZIP_DEFLATED) as z:
            for f in target.rglob('*'):
                if cancel():raise InterruptedError('打包已取消，结果未完成。')
                if f.is_file() and f not in (zip_path,target/'导出记录.json'): z.write(f,f.relative_to(target).as_posix())
        with zipfile.ZipFile(zip_path) as z:
            if z.testzip(): raise OSError('ZIP完整性检查失败')
            for m,relative in plan:
                member=relative.as_posix()
                if member not in z.namelist() or z.getinfo(member).file_size!=m.size:raise OSError('ZIP成员缺失或大小不符：'+member)
                h=hashlib.sha256()
                with z.open(member) as f:
                    for block in iter(lambda:f.read(1024*1024),b''):h.update(block)
                if h.hexdigest()!=m.sha256:raise OSError('ZIP成员内容校验失败：'+member)
        state['status']='完成';manifest()
        project.exports.append({'time':now(),'path':str(target),'status':'完成'}); return target
    except Exception as e:
        zip_path=target/'整理结果.zip'
        if zip_path.exists():zip_path.rename(target/'整理结果.zip.partial')
        state['status']='未完成'; state['error']=str(e); manifest(); raise type(e)(str(e)+'\n未完成目录：'+str(target)+'\n请处理原因后重新导出，不要交付此目录。') from e
