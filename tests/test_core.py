import sys, csv, copy, zipfile, json, time
from pathlib import Path
sys.path.insert(0,str(Path(__file__).parents[1]/'src'))
from core import *
import pytest

@pytest.mark.parametrize('filename,people,expected',[
 ('0012_照片.jpg',[('张三','0012')],0),('10012_照片.jpg',[('张三','0012')],None),
 ('00120.jpg',[('张三','0012')],None),('a0012.jpg',[('张三','0012')],None),
 ('0012a.jpg',[('张三','0012')],None),('(0012).jpg',[('张三','0012')],0),
 ('0012-照片.jpg',[('张三','0012')],0),('0012 照片.jpg',[('张三','0012')],0),
 ('0012_10012.jpg',[('甲','0012'),('乙','10012')],None),
 ('0012.jpg',[('甲','0012'),('乙','0012')],None),
 ('张三.jpg',[('张三','1')],0),('张三丰.jpg',[('张三','1'),('张三丰','2')],1),
 ('张三丰.jpg',[('张三','1')],None),('张三.jpg',[('张三','1'),('张三','2')],None),
 ('张三_照片.jpg',[('张三','1')],0),('我是张三.jpg',[('张三','1')],None),
 ('张三同学.jpg',[('张三','1')],None),('张三-0012.jpg',[('张三','0012')],0),
 ('李四_0012.jpg',[('张三','0012'),('李四','0013')],None),
 ('匿名.jpg',[('张三','0012')],None),('00012.jpg',[('张三','0012')],None),
 ('12.jpg',[('张三','0012')],None),('0012.pdf.jpg',[('张三','0012')],0),
 ('ABC-12.jpg',[('甲','ABC-12')],0),('XABC-12.jpg',[('甲','ABC-12')],None),
 ('abc-12.jpg',[('甲','ABC-12')],0),('0012（照片）.jpg',[('甲','0012')],0),
 ('张 三.jpg',[('张三','1')],None),('王小明.jpg',[('王小','1')],None),
 ('张三_李四.jpg',[('张三','1'),('李四','2')],None),('编号0012.jpg',[('甲','0012')],None),
 ('0012_张三丰.jpg',[('张三','0012'),('张三丰','2')],None),
])
def test_match_boundaries(filename,people,expected):
    p=Project();p.people=[Person(str(i),n,id) for i,(n,id) in enumerate(people)];p.rules=[Rule('照片',['.jpg'])]
    m=Material(filename,filename,0,'');match(p,m)
    assert m.person==(None if expected is None else str(expected))

def test_roster_and_encoding(tmp_path):
    p=tmp_path/'名单.csv';p.write_bytes('姓名,学号\n张三,0012\n\n李四,10012'.encode('gb18030'))
    h,r,w=read_roster(p);people,_=parse_people(h,r,0,1);assert [p.identifier for p in people]==['0012','10012']
    people,w=parse_people(['姓名','手机号后四位'],[['甲','1234']],0,1);assert people[0].identifier=='' and w

def test_xlsx(tmp_path):
    from openpyxl import Workbook
    p=tmp_path/'名单.xlsx';wb=Workbook();ws=wb.active;ws.append(['姓名','学号']);ws.append(['甲',12]);ws['B2'].number_format='0000';wb.save(p)
    h,r,w=read_roster(p);assert r[0][1]=='0012'

@pytest.mark.parametrize('count',[10,100,1000])
def test_end_to_end(tmp_path,count):
    source=tmp_path/'源材料';source.mkdir();dest=tmp_path/'交付';p=Project();p.rules=[Rule('作品',['.pdf'])]
    for i in range(count):
        p.people.append(Person(str(i),f'虚构姓名{i}',f'{i:06}'))
        if i%10:(source/f'{i:06}_作品.pdf').write_bytes(f'fictional {i}'.encode())
    (source/'匿名.pdf').write_bytes(b'unknown')
    before={str(f):digest(f) for f in source.iterdir()};start=time.perf_counter();files,errors=scan(source);merge_scan(p,source,files)
    assert not errors and sum(bool(status(p,x)[1]) for x in p.people)==count//10
    anon=next(i for i,m in enumerate(p.materials) if '匿名' in m.relative);assign(p,anon,'0','作品')
    project=tmp_path/'project.msq';save(p,project);q=load(project);assert q.materials[anon].manual
    result=export(q,dest)
    from openpyxl import load_workbook
    for name in ['核对总表.xlsx','缺件清单.xlsx','异常文件清单.xlsx']:
        wb=load_workbook(result/name);assert wb.active.max_row>1;wb.close()
    with zipfile.ZipFile(result/'整理结果.zip') as z:assert z.testzip() is None;z.extractall(tmp_path/'验包')
    assert before=={str(f):digest(f) for f in source.iterdir()}
    print(f'case={count}, files={len(files)}, seconds={time.perf_counter()-start:.3f}')

def test_changed_manual_and_cancel(tmp_path):
    root=tmp_path/'源';root.mkdir();f=root/'匿名.pdf';f.write_bytes(b'1');p=Project(people=[Person('a','甲')],rules=[Rule('作品',['.pdf'])]);files,_=scan(root);merge_scan(p,root,files);assign(p,0,'a','作品')
    files,_=scan(root);merge_scan(p,root,files);assert p.materials[0].person=='a'
    f.write_bytes(b'2');files,_=scan(root);merge_scan(p,root,files);assert p.materials[0].needs_review;assert status(p,p.people[0])[1]
    with pytest.raises(InterruptedError):scan(root,lambda:True)
    with pytest.raises(InterruptedError):export(p,tmp_path/'输出',lambda:True)

def test_overlap_and_missing(tmp_path):
    root=tmp_path/'源';root.mkdir();f=root/'001.pdf';f.write_bytes(b'1');p=Project(people=[Person('a','甲','001')],rules=[Rule('作品',['.pdf'])]);files,_=scan(root);merge_scan(p,root,files)
    for out in [root,root/'sub',tmp_path]:
        with pytest.raises(ValueError):export_plan(p,out)
    f.unlink()
    with pytest.raises(ValueError):export(p,tmp_path/'输出')

def test_formula_and_duplicates(tmp_path):
    root=tmp_path/'源';root.mkdir();(root/'001_a.pdf').write_bytes(b'a');(root/'001_b.pdf').write_bytes(b'b')
    p=Project(people=[Person('a','=HYPERLINK("evil")','001')],rules=[Rule('作品',['.pdf'])]);files,_=scan(root);merge_scan(p,root,files);assert status(p,p.people[0])[2]==['作品']
    result=export(p,tmp_path/'输出');from openpyxl import load_workbook
    wb=load_workbook(result/'核对总表.xlsx');assert wb.active['A2'].data_type=='s';wb.close()
    assert len(list((result/'材料').rglob('*.pdf')))==2

@pytest.mark.parametrize('name',['../escape','CON','LPT1.txt','a/b','x\\y','a:bad','..','NUL','bad\x00x'])
def test_safe_names(name):
    s=safe_name(name);assert '/' not in s and '\\' not in s and ':' not in s and s not in ('.','..')

def test_corrupt_and_backup(tmp_path):
    f=tmp_path/'a.msq';p=Project();save(p,f);p.title='第二版';save(p,f);assert load(str(f)+'.bak').title=='新核对任务';f.write_text('{broken')
    with pytest.raises(json.JSONDecodeError):load(f)

def test_backup_hardlink_source_protection(tmp_path):
    f=tmp_path/'source.txt';f.write_bytes(b'original');dest=tmp_path/'project.msq';p=Project();save(p,dest)
    os.link(f,str(dest)+'.bak')
    with pytest.raises(ValueError):save(p,dest)
    assert f.read_bytes()==b'original'

def test_save_source_and_manual_format(tmp_path):
    root=tmp_path/'源';root.mkdir();f=root/'001.jpg';f.write_bytes(b'image');p=Project(people=[Person('a','甲','001')],rules=[Rule('作品',['.pdf'])]);files,_=scan(root);merge_scan(p,root,files)
    with pytest.raises(ValueError):save(p,f)
    with pytest.raises(ValueError):assign(p,0,'a','作品')
    assert f.read_bytes()==b'image'

def test_file_lock_on_windows(tmp_path):
    import ctypes
    if os.name!='nt':pytest.skip('Windows only')
    f=tmp_path/'锁定.pdf';f.write_bytes(b'lock')
    kernel=ctypes.WinDLL('kernel32',use_last_error=True)
    kernel.CreateFileW.argtypes=[ctypes.c_wchar_p,ctypes.c_uint32,ctypes.c_uint32,ctypes.c_void_p,ctypes.c_uint32,ctypes.c_uint32,ctypes.c_void_p];kernel.CreateFileW.restype=ctypes.c_void_p
    handle=kernel.CreateFileW(str(f),0x80000000,0,None,3,0,None)
    if handle==ctypes.c_void_p(-1).value:pytest.skip('Cannot acquire exclusive synthetic file handle')
    try:
        files,errors=scan(tmp_path);assert not files and errors
    finally:kernel.CloseHandle.argtypes=[ctypes.c_void_p];kernel.CloseHandle(handle)
    assert f.read_bytes()==b'lock'

def test_junction_rejected(tmp_path):
    import subprocess
    if os.name!='nt':pytest.skip('NTFS only')
    real=tmp_path/'real';real.mkdir();(real/'a.pdf').write_bytes(b'a');link=tmp_path/'junction'
    result=subprocess.run(['cmd','/c','mklink','/J',str(link),str(real)],capture_output=True)
    if result.returncode:pytest.skip('Junction creation unavailable')
    with pytest.raises(ValueError):scan(link)
    with pytest.raises(ValueError):export_plan(Project(),link/'output')

def test_empty_bad_headers(tmp_path):
    f=tmp_path/'empty.csv';f.write_text('姓名,学号\n',encoding='utf-8')
    with pytest.raises(ValueError):read_roster(f)
    with pytest.raises(ValueError):parse_people(['姓名'],[['']],0)
    with pytest.raises(ValueError):parse_people(['姓名'],[['甲']],2)

def test_deleted_changed_status(tmp_path):
    f=tmp_path/'001.pdf';f.write_bytes(b'1');p=Project(people=[Person('a','甲','001')],rules=[Rule('作品',['.pdf'])]);files,_=scan(tmp_path);merge_scan(p,tmp_path,files);assert not status(p,p.people[0])[1]
    f.write_bytes(b'changed');assert status(p,p.people[0])[1]
    f.unlink();files,_=scan(tmp_path);merge_scan(p,tmp_path,files);assert p.materials[0].invalid;assert status(p,p.people[0])[1]

def test_ambiguous_categories(tmp_path):
    p=Project(people=[Person('a','甲','001')],rules=[Rule('报名',['.pdf'],keywords=['报名']),Rule('作品',['.pdf'],keywords=['作品'])]);m=Material('001_报名_作品.pdf','001_报名_作品.pdf',0,'');match(p,m);assert m.category is None

def test_untrusted_project_key_and_headers(tmp_path):
    from openpyxl import load_workbook
    root=tmp_path/'source';root.mkdir();(root/'001.pdf').write_bytes(b'pdf')
    p=Project(project_id='=1+1',people=[Person('x/../../','甲','001')],rules=[Rule('=1+1',['.pdf'])])
    files,_=scan(root);merge_scan(p,root,files)
    for _,relative in export_plan(p,tmp_path/'output'):assert '..' not in relative.parts
    result=export(p,tmp_path/'output');wb=load_workbook(result/'核对总表.xlsx')
    assert all(c.data_type!='f' for row in wb.active for c in row);wb.close()

def test_low_space_preflight_unit_injection(tmp_path,monkeypatch):
    import core
    from collections import namedtuple
    usage=namedtuple('Usage','total used free')
    monkeypatch.setattr(core.shutil,'disk_usage',lambda p:usage(100,99,1))
    with pytest.raises(OSError):export(Project(),tmp_path/'output')

def test_save_never_overwrites_unrelated_original(tmp_path):
    f=tmp_path/'名单.csv';f.write_bytes(b'original-roster')
    with pytest.raises(ValueError):save(Project(),f)
    assert f.read_bytes()==b'original-roster'

def test_damaged_excel_is_explained_in_chinese(tmp_path):
    f=tmp_path/'坏名单.xlsx';f.write_bytes(b'not xlsx')
    with pytest.raises(ValueError,match='Excel 名单无法读取'):read_roster(f)

def test_unrelated_backup_never_overwritten(tmp_path):
    f=tmp_path/'project.msq';save(Project(),f);backup=Path(str(f)+'.bak');backup.write_bytes(b'unrelated')
    with pytest.raises(ValueError):save(Project(),f)
    assert backup.read_bytes()==b'unrelated'

def test_zip_failure_is_partial_unit_injection(tmp_path,monkeypatch):
    import core
    original=core.zipfile.ZipFile.testzip
    monkeypatch.setattr(core.zipfile.ZipFile,'testzip',lambda self:'broken-member')
    with pytest.raises(OSError):export(Project(),tmp_path/'output')
    result=next((tmp_path/'output').iterdir());assert not (result/'整理结果.zip').exists();assert (result/'整理结果.zip.partial').exists()
    assert json.loads((result/'导出记录.json').read_text(encoding='utf-8'))['status']=='未完成'
    monkeypatch.setattr(core.zipfile.ZipFile,'testzip',original)

def test_other_valid_project_backup_is_preserved(tmp_path):
    f=tmp_path/'current.msq';p=Project();save(p,f);other=Project();b=Path(str(f)+'.bak');save(other,b);before=digest(b)
    with pytest.raises(ValueError):save(p,f)
    assert digest(b)==before and load(b).project_id==other.project_id

def test_real_ntfs_replacement_same_bytes_requires_review(tmp_path):
    f=tmp_path/'unknown.pdf';f.write_bytes(b'same');p=Project(people=[Person('a','甲')],rules=[Rule('文件',['.pdf'])]);found,_=scan(tmp_path);merge_scan(p,tmp_path,found);assign(p,0,'a','文件')
    replacement=tmp_path/'replacement.pdf';replacement.write_bytes(b'same');old=f.stat().st_ino;os.replace(replacement,f);assert f.stat().st_ino!=old
    found,_=scan(tmp_path);merge_scan(p,tmp_path,found);assert p.materials[0].needs_review and status(p,p.people[0])[1]

def test_real_ntfs_ads_is_rejected(tmp_path):
    if os.name!='nt':pytest.skip('NTFS ADS only')
    f=tmp_path/'source.pdf';f.write_bytes(b'base');ads=Path(str(f)+':secret');ads.write_bytes(b'secret');s=ads.stat();p=Project(materials=[Material(str(ads),'source.pdf:secret',s.st_size,digest(ads),mtime_ns=s.st_mtime_ns)])
    jf=tmp_path/'bad.msq';save(p,jf)
    with pytest.raises(ValueError):load(jf)
    with pytest.raises(ValueError):export(p,tmp_path/'output')
    assert ads.read_bytes()==b'secret'

def test_rules_preserve_unaffected_manual_and_recover_affected(tmp_path):
    (tmp_path/'a.pdf').write_bytes(b'a');(tmp_path/'b.png').write_bytes(b'b')
    p=Project(people=[Person('a','甲')],rules=[Rule('文件',['.pdf']),Rule('照片',['.png'])]);found,_=scan(tmp_path);merge_scan(p,tmp_path,found)
    for i,m in enumerate(p.materials):assign(p,i,'a','文件' if m.path.endswith('.pdf') else '照片')
    changed=copy.deepcopy(p.rules);changed[0].keywords=['申请'];assert update_rules(p,changed)==1
    pdf=next(m for m in p.materials if m.path.endswith('.pdf'));photo=next(m for m in p.materials if m.path.endswith('.png'))
    assert pdf.person=='a' and pdf.needs_review and photo.person=='a' and not photo.needs_review
    save(p,tmp_path/'project.msq');q=load(tmp_path/'project.msq');assert sum(m.needs_review for m in q.materials)==1
    assign(p,p.materials.index(pdf),'a','文件');assert not pdf.needs_review

def test_missing_is_persisted_and_explicit_ignore_keeps_missing_report(tmp_path):
    source=tmp_path/'source';source.mkdir();f=source/'001.pdf';f.write_bytes(b'file');p=Project(people=[Person('a','甲','001')],rules=[Rule('文件',['.pdf'])]);found,_=scan(source);merge_scan(p,source,found);f.unlink();found,_=scan(source);merge_scan(p,source,found)
    save(p,tmp_path/'missing.msq');q=load(tmp_path/'missing.msq');assert q.materials[0].invalid and status(q,q.people[0])[1]
    with pytest.raises(ValueError,match='重新扫描'):export_plan(q,tmp_path/'output')
    q.materials[0].ignored=True;result=export(q,tmp_path/'output')
    from openpyxl import load_workbook
    wb=load_workbook(result/'缺件清单.xlsx');assert wb.active['A2'].value=='甲';wb.close()

@pytest.mark.parametrize('filename,name,identifier',[('JoAnn.pdf','Ann',''),('１0012.pdf','甲','0012'),('①0012.pdf','甲','0012')])
def test_unicode_identity_boundary(filename,name,identifier):
    p=Project(people=[Person('a',name,identifier)],rules=[Rule('文件',['.pdf'])]);m=Material(filename,filename,0,'');match(p,m);assert m.person is None

def test_unicode_long_path_actual_filesystem(tmp_path):
    root=tmp_path/'source';directory=root
    for i in range(6):directory=directory/('中文 空格（材料）'+str(i)*32)
    try:directory.mkdir(parents=True);f=directory/'001_作品.pdf';f.write_bytes(b'long-path')
    except OSError as e:pytest.skip('Actual filesystem long path unsupported: '+str(e))
    files,errors=scan(root);assert len(files)==1 and not errors
    p=Project(people=[Person('a','甲','001')],rules=[Rule('作品',['.pdf'])]);merge_scan(p,root,files)
    result=export(p,tmp_path/'output');assert f.read_bytes()==b'long-path';assert list((result/'材料').rglob('*.pdf'))

@pytest.mark.parametrize('filename,name,identifier',[('Jose\u0301.pdf','Jose',''),('Ann\u0301.pdf','Ann',''),('0012\u0301.pdf','甲','0012'),('\u03010012.pdf','甲','0012')])
def test_combining_mark_is_not_identity_separator(filename,name,identifier):
    p=Project(people=[Person('a',name,identifier)],rules=[Rule('文件',['.pdf'])]);m=Material(filename,filename,0,'');match(p,m);assert m.person is None

def test_ignored_real_ntfs_replacement_requires_review(tmp_path):
    f=tmp_path/'0012.pdf';f.write_bytes(b'same');p=Project(people=[Person('a','甲','0012')],rules=[Rule('文件',['.pdf'])]);found,_=scan(tmp_path);merge_scan(p,tmp_path,found);p.materials[0].ignored=True
    replacement=tmp_path/'replacement.pdf';replacement.write_bytes(b'same');old=f.stat().st_ino;os.replace(replacement,f);assert f.stat().st_ino!=old
    found,_=scan(tmp_path);merge_scan(p,tmp_path,found);assert p.materials[0].needs_review and not p.materials[0].ignored and p.materials[0].person is None

def test_legacy_ignored_without_file_identity_requires_review(tmp_path):
    f=tmp_path/'0012.pdf';f.write_bytes(b'file');p=Project(people=[Person('a','甲','0012')],rules=[Rule('文件',['.pdf'])]);found,_=scan(tmp_path);merge_scan(p,tmp_path,found);p.materials[0].ignored=True;p.materials[0].file_id=''
    save(p,tmp_path/'legacy.msq');q=load(tmp_path/'legacy.msq');assert q.materials[0].needs_review and not q.materials[0].ignored
    with pytest.raises(ValueError,match='待复核'):export_plan(q,tmp_path.parent/(tmp_path.name+'_output'))
